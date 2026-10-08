"""Management acceptance tests use only disposable in-memory databases and temp files."""
import io
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path

from flask import Flask, g
from flask_login import current_user
from PIL import Image
from werkzeug.datastructures import FileStorage

from models import (Organization, OrganizationWorkspace, Department, User, FinancialRecord,
    FinanceLedgerEntry, FinanceAccount, FinanceReceivable, FinancePayable,
    BudgetAllocation, BudgetReservation, ApprovalRule, AuditRevision, InventoryItem,
    StockMovement, Attachment, ManagementTask, Notification, ReconciliationMatch)
from packages.database import db
from packages.security import init_security
from modules.control_routes import register_control_routes
from services.audit_service import install_audit_tracking
from services.department_operations import create_department_operation
from services.finance_service import create_reconciliation, post_ledger_entry, settle_receivable, settle_payable
from services.operations_service import create_request, decide_request, save_fulfillment, move_stock
from services.owner_control_service import management_workspace, scoped_finance_summary
from services.record_trace_service import record_trace
from services.reporting_service import management_export_rows
from services.period_service import period_bounds
from services.export_service import build_report_export


class OwnerControlTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="sage-owner-tests-")
        self.app = Flask(__name__)
        self.app.config.update(SECRET_KEY="test-only-secret", SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",
            TESTING=True, STORAGE_DIR=Path(self.temp.name), MAX_UPLOAD_MB=12,
            EMAIL_NOTIFICATIONS_ENABLED=False, WEB_PUSH_ENABLED=False)
        db.init_app(self.app); init_security(self.app); register_control_routes(self.app)
        install_audit_tracking(self.app)
        @self.app.before_request
        def actor():
            # Each production request has a fresh g; the outer test app context must emulate it.
            g.pop("_login_user", None)
            db.session.info["sage_actor_id"] = current_user.id if current_user.is_authenticated else None
        self.context = self.app.app_context(); self.context.push(); db.create_all()
        self.org = Organization(name="Acceptance fixture business", slug="acceptance", business_type="Retail")
        self.other_org = Organization(name="Isolated fixture", slug="isolated", business_type="Retail")
        db.session.add_all([self.org,self.other_org]); db.session.flush()
        self.dept = Department(organization_id=self.org.id,name="Sales")
        self.private_dept = Department(organization_id=self.org.id,name="Finance")
        self.other_dept = Department(organization_id=self.other_org.id,name="Sales")
        db.session.add_all([self.dept,self.private_dept,self.other_dept]); db.session.flush()
        def worker(name,role,org,dept):
            row = User(organization_id=org.id,department_id=dept.id if dept else None,first_name=name,last_name="Fixture",email=name.lower()+"@example.test",password_hash="not-a-real-password",role=role)
            db.session.add(row); return row
        self.owner = worker("Owner","owner",self.org,self.dept)
        self.staff = worker("Staff","staff",self.org,self.dept)
        self.head = worker("Head","department_head",self.org,self.dept)
        self.finance = worker("Finance","finance",self.org,self.private_dept)
        self.unassigned = worker("Unassigned","staff",self.org,None)
        self.outsider = worker("Outside","owner",self.other_org,self.other_dept)
        db.session.flush()
        self.now = datetime.now(timezone.utc)
        self.bank = FinanceAccount(organization_id=self.org.id,name="Fixture bank",opening_balance=500,created_by_id=self.owner.id)
        self.budget = BudgetAllocation(organization_id=self.org.id,department_id=self.dept.id,created_by_id=self.owner.id,name="Fixture cost centre",original_budget=2000,revised_budget=2000,period_start=self.now.date()-timedelta(days=30),period_end=self.now.date()+timedelta(days=30))
        db.session.add_all([self.bank,self.budget]); db.session.commit()

    def tearDown(self):
        db.session.remove(); db.engine.dispose(); self.context.pop(); self.temp.cleanup()

    def client(self,user):
        client = self.app.test_client()
        with client.session_transaction() as session: session['_user_id']=user.id; session['_fresh']=True
        return client

    def receipt(self,name="fixture-receipt.png",color="white"):
        output=io.BytesIO(); Image.new("RGB",(20,20),color).save(output,"PNG"); output.seek(0)
        return FileStorage(output,filename=name,content_type="image/png")

    def operation(self,kind,amount,user=None):
        user=user or self.staff; db.session.info['sage_actor_id']=user.id
        return create_department_operation(self.app,user,{"operation_type":kind,"amount":amount,"title":"Acceptance test actual financial activity"})

    def test_empty_records_do_not_invent_metrics(self):
        data=management_workspace(self.owner)
        self.assertEqual(data['totals']['income'],0); self.assertEqual(data['totals']['expenses'],0)
        self.assertIsNone(data['insights']['evidence_coverage']); self.assertIsNone(data['insights']['purchase_cycle_days'])
        self.assertIsNone(data['insights']['cash_coverage_days'])

    def test_staff_activity_posts_one_financial_record_and_notifies_owner(self):
        income=self.operation('financial_income','125.50'); expense=self.operation('financial_expense','35.25')
        self.assertEqual(scoped_finance_summary(self.staff)['net'],90.25)
        self.assertEqual(FinancialRecord.query.filter_by(source_entity_id=income.id).count(),1)
        self.assertEqual(FinanceLedgerEntry.query.filter_by(source_entity_id=income.id).count(),1)
        self.assertEqual(float(self.budget.actual_spend),35.25)
        self.assertGreater(Notification.query.filter_by(user_id=self.owner.id).count(),0)
        data=management_workspace(self.owner,'all'); self.assertEqual(data['totals']['cash_in'],125.50)
        self.assertEqual(data['totals']['cash_out'],35.25)
        self.assertIn(expense.id,[row['source_entity_id'] for row in [dict(source_entity_id=item.source_entity_id) for item in FinancialRecord.query.all()]])

    def test_unassigned_staff_cannot_see_company_finances(self):
        self.operation('financial_income',200)
        self.assertEqual(scoped_finance_summary(self.unassigned)['income'],0)
        self.assertEqual(self.client(self.unassigned).get('/api/management/workspace').status_code,403)

    def test_private_finance_posting_is_not_exposed_to_ordinary_staff(self):
        row=post_ledger_entry(self.app,self.owner,{'entry_type':'expense','amount':'45','description':'Private finance fixture','department_id':self.dept.id})
        self.assertEqual(self.client(self.staff).get('/api/management/records/finance_ledger/'+row.id).status_code,404)

    def test_cross_tenant_history_and_links_are_rejected(self):
        row=FinancialRecord(reference="ISOLATED",organization_id=self.other_org.id,created_by_id=self.outsider.id,description="Isolated fixture",record_type="revenue",amount=999)
        db.session.add(row); db.session.commit()
        self.assertEqual(self.client(self.owner).get('/api/management/records/financial_record/'+row.id).status_code,404)
        self.assertEqual(management_workspace(self.owner,'all')['totals']['income'],0)

    def test_non_financial_activity_does_not_fabricate_income(self):
        self.operation('activity_completed',999)
        self.assertEqual(FinancialRecord.query.count(),0)

    def test_cash_movement_is_not_a_second_expense(self):
        post_ledger_entry(self.app,self.owner,{'entry_type':'cash_out','amount':100,'description':'Cash advance fixture','account_id':self.bank.id})
        data=management_workspace(self.owner,'all')
        self.assertEqual(data['totals']['expenses'],0); self.assertEqual(data['totals']['cash_out'],100)
        self.assertEqual(data['totals']['available_cash'],400)

    def test_budget_reservation_becomes_actual_without_double_counting(self):
        purchase=create_request(self.app,self.staff,{'department_id':self.dept.id,'item_name':'Fixture supplies','quantity':2,'unit_cost':50})
        db.session.add(ApprovalRule(organization_id=self.org.id,name="Small fixture purchases",min_amount=0,max_amount=500,required_role="department_head"));db.session.commit()
        decide_request(self.app,self.head,purchase,'approve')
        self.assertEqual(float(self.budget.committed_amount),100)
        record=save_fulfillment(self.app,self.staff,purchase,{'finalize':'true','supplier_name':'Fixture vendor','actual_quantity_'+purchase.items[0].id:'2','actual_unit_cost_'+purchase.items[0].id:'45'},self.receipt())
        self.assertEqual(float(self.budget.committed_amount),0); self.assertEqual(float(self.budget.actual_spend),90)
        self.assertEqual(BudgetReservation.query.first().status,'spent')
        trace=record_trace(self.owner,'request',purchase.id)
        self.assertTrue(any(row['type']=='inventory_item' for row in trace['records']))
        self.assertEqual(len(trace['evidence']),1)
        self.assertFalse(next(stage for stage in trace['stages'] if stage['label']=='Payment')['done'])
        self.assertEqual(management_workspace(self.owner,'all')['insights']['evidence_coverage'],100)

    def test_rejection_releases_budget(self):
        purchase=create_request(self.app,self.staff,{'item_name':'Fixture purchase','quantity':1,'unit_cost':120})
        decide_request(self.app,self.owner,purchase,'approve'); decide_request(self.app,self.owner,purchase,'reject')
        self.assertEqual(float(self.budget.committed_amount),0)

    def test_self_approval_and_uncovered_limits_do_not_bypass_controls(self):
        purchase=create_request(self.app,self.head,{'item_name':'Fixture purchase','quantity':1,'unit_cost':150})
        db.session.add(ApprovalRule(organization_id=self.org.id,name="Configured fixture range",min_amount=0,max_amount=100,required_role="department_head"));db.session.commit()
        with self.assertRaises(PermissionError): decide_request(self.app,self.head,purchase,'approve')
        with self.assertRaises(PermissionError): decide_request(self.app,self.finance,purchase,'approve')
        decide_request(self.app,self.owner,purchase,'approve')

    def test_historical_reconciliation_excludes_future_entries_and_matches_selected(self):
        old=post_ledger_entry(self.app,self.owner,{'entry_type':'expense','amount':100,'description':'Historical payment','account_id':self.bank.id,'occurred_at':(self.now-timedelta(days=10)).isoformat()})
        post_ledger_entry(self.app,self.owner,{'entry_type':'revenue','amount':300,'description':'Later receipt','account_id':self.bank.id,'occurred_at':self.now.isoformat()})
        reconciliation=create_reconciliation(self.app,self.owner,{'account_id':self.bank.id,'period_start':(self.now-timedelta(days=20)).date().isoformat(),'period_end':(self.now-timedelta(days=5)).date().isoformat(),'statement_balance':'400','ledger_ids':[old.id]})
        self.assertEqual(float(reconciliation.system_balance),400);self.assertEqual(old.status,'reconciled')
        self.assertEqual(ReconciliationMatch.query.count(),1)

    def test_reconciliation_variance_does_not_mark_unmatched_entries(self):
        entry=post_ledger_entry(self.app,self.owner,{'entry_type':'expense','amount':100,'description':'Fixture payment','account_id':self.bank.id})
        with self.assertRaises(ValueError):create_reconciliation(self.app,self.owner,{'account_id':self.bank.id,'period_start':self.now.date().isoformat(),'period_end':self.now.date().isoformat(),'statement_balance':399,'ledger_ids':[entry.id]})
        db.session.rollback();self.assertEqual(entry.status,'posted');self.assertEqual(ReconciliationMatch.query.count(),0)

    def test_aging_and_payment_history_use_recorded_dates_and_amounts(self):
        ar=FinanceReceivable(reference="AR-FIXTURE",organization_id=self.org.id,customer_name="Fixture customer",description="Fixture receivable",total_amount=200,amount_received=0,due_date=self.now.date()-timedelta(days=45),created_by_id=self.owner.id)
        ap=FinancePayable(reference="AP-FIXTURE",organization_id=self.org.id,vendor_name="Fixture vendor",description="Fixture payable",total_amount=100,amount_paid=0,due_date=self.now.date()+timedelta(days=3),created_by_id=self.owner.id)
        db.session.add_all([ar,ap]);db.session.commit()
        settle_receivable(self.app,self.owner,ar,25);settle_payable(self.app,self.owner,ap,20)
        data=management_workspace(self.owner,'all')
        self.assertEqual(data['receivables']['buckets']['31–60 days'],175)
        self.assertEqual(data['receivables']['rows'][0]['payments'][0]['amount'],'25.00')
        self.assertEqual(data['payables']['total'],80);self.assertEqual(len(data['insights']['due_soon']),1)

    def test_staff_can_complete_only_their_assigned_followup(self):
        response=self.client(self.owner).post('/api/management/tasks',json={'title':'Fixture follow-up','assigned_to_id':self.staff.id,'due_date':self.now.date().isoformat()})
        self.assertEqual(response.status_code,200);task_id=response.json['task']['id']
        response=self.client(self.staff).patch('/api/management/tasks/'+task_id,json={'status':'completed','comment':'Completed fixture work'})
        self.assertEqual(response.status_code,200);self.assertEqual(response.json['task']['status'],'completed')
        self.assertEqual(len(response.json['task']['comments']),1)
        self.assertEqual(self.client(self.head).patch('/api/management/tasks/'+task_id,json={'status':'open'}).status_code,403)

    def test_expense_verification_requires_evidence_and_owner_attestation_is_explicit(self):
        row=post_ledger_entry(self.app,self.owner,{'entry_type':'expense','amount':25,'description':'Owner fixture expense'})
        client=self.client(self.owner)
        self.assertEqual(client.post('/api/management/records/finance_ledger/'+row.id+'/verify',json={'notes':'Checked statement'}).status_code,400)
        response=client.post('/api/management/records/finance_ledger/'+row.id+'/evidence',data={'evidence':(self.receipt().stream,'fixture-receipt.png')},content_type='multipart/form-data')
        self.assertEqual(response.status_code,200)
        self.assertEqual(client.post('/api/management/records/finance_ledger/'+row.id+'/verify',json={'notes':'Checked receipt and payment'}).status_code,200)
        history=record_trace(self.owner,'finance_ledger',row.id)
        self.assertTrue(any('Owner attestation:' in item.get('description','') for item in history['history']))

    def test_deleted_record_retains_full_before_snapshot(self):
        db.session.info['sage_actor_id']=self.owner.id
        item=InventoryItem(organization_id=self.org.id,department_id=self.dept.id,name='Retained fixture',quantity=3,unit_value=15,added_by_id=self.owner.id)
        db.session.add(item);db.session.commit();item_id=item.id
        db.session.delete(item);db.session.commit()
        history=record_trace(self.owner,'inventory_item',item_id)
        self.assertTrue(history['record']['archived'])
        self.assertTrue(any(row.get('before',{}).get('name')=='Retained fixture' for row in history['history']))

    def test_audit_history_cannot_be_rewritten(self):
        row=AuditRevision.query.first();row.change_summary='Changed history'
        with self.assertRaises(ValueError):db.session.commit()
        db.session.rollback()

    def test_exports_keep_numbers_and_escape_untrusted_formula_labels(self):
        from openpyxl import load_workbook
        payload,_=build_report_export(['Description','Amount'],[['=1+1','-25.50']],'xlsx','Fixture export')
        sheet=load_workbook(payload).active
        self.assertEqual(sheet.cell(6,1).data_type,'s');self.assertEqual(sheet.cell(6,2).value,-25.5)
        payload,_=build_report_export(['Description','Amount'],[['=1+1','-25.50']],'csv','Fixture export')
        self.assertIn("'=1+1",payload.getvalue().decode('utf-8-sig'))

    def test_period_end_excludes_future_period_records(self):
        start,end=period_bounds('monthly')
        db.session.add(FinancialRecord(reference="FUTURE",organization_id=self.org.id,record_type="revenue",description="Future fixture",amount=999,occurred_at=end));db.session.commit()
        self.assertEqual(management_workspace(self.owner,'monthly')['totals']['income'],0)
        headers,rows=management_export_rows(self.owner,'income','monthly');self.assertEqual(rows,[])

    def test_individual_business_uses_same_real_management_data(self):
        db.session.add(OrganizationWorkspace(organization_id=self.org.id,workspace_mode='individual'));db.session.commit()
        self.operation('financial_income',77)
        self.assertEqual(self.client(self.owner).get('/api/management/workspace?period=all').json['data']['totals']['income'],77)

    def test_direct_stock_purchase_posts_cost_once_and_links_supplier_prices(self):
        row=post_ledger_entry(self.app,self.owner,{'entry_type':'expense','amount':'90','description':'Direct fixture purchase','counterparty':'Fixture vendor','department_id':self.dept.id,'record_as':'stock','purchase_item_name':'Fixture direct supplies','purchase_quantity':'2','purchase_unit_cost':'45','purchase_unit':'unit'},self.receipt())
        self.assertEqual(InventoryItem.query.first().total_value,90)
        self.assertEqual(FinancialRecord.query.count(),1)
        trace=record_trace(self.owner,'finance_ledger',row.id)
        self.assertEqual(trace['purchase_lines'][0]['unit_cost'],45)
        self.assertTrue(any(item['type']=='inventory_item' for item in trace['records']))
        self.assertEqual(management_workspace(self.owner,'all')['suppliers'][0]['prices'][0]['price'],45)

    def test_actual_invoice_date_controls_owing_age(self):
        row=FinanceReceivable(reference="DATE-FIXTURE",organization_id=self.org.id,customer_name="Fixture customer",description="Invoice fixture",total_amount=100,amount_received=0,created_by_id=self.owner.id)
        db.session.add(row);db.session.commit()
        issued=(self.now-timedelta(days=60)).date().isoformat()
        response=self.client(self.owner).patch('/api/management/obligations/receivable/'+row.id+'/terms',json={'invoice_date':issued})
        self.assertEqual(response.status_code,200)
        self.assertEqual(management_workspace(self.owner,'all')['receivables']['rows'][0]['age_days'],60)

    def test_named_recipient_can_acknowledge_cross_department_stock_issue(self):
        item=InventoryItem(organization_id=self.org.id,department_id=self.private_dept.id,name='Issued fixture',quantity=3,unit_value=10,added_by_id=self.owner.id)
        db.session.add(item);db.session.commit()
        movement=move_stock(self.app,self.owner,item,'stock_out',1,recipient=self.staff)
        response=self.client(self.staff).post('/api/management/records/stock_movement/'+movement.id+'/acknowledge',json={'action':'received','notes':'Received the fixture stock myself'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.client(self.head).post('/api/management/records/stock_movement/'+movement.id+'/acknowledge',json={'action':'received','notes':'Not my assigned stock'}).status_code,404)


if __name__=='__main__':unittest.main()
