"""Shared, safe CSV/Excel exports and readable PDF tables for real saved records."""
import csv
import io
import re
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

def _cell(value): return "" if value is None else str(value)

def _spreadsheet_cell(value):
    text = _cell(value)
    if text.lstrip().startswith(("=", "+", "@", "-")) and not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text): return "'" + text
    return text

def build_report_export(headers, rows, file_format, title, subtitle=""):
    file_format = str(file_format or "csv").strip().lower()
    headers = [_cell(value) for value in headers]
    safe_rows = [[_cell(value) for value in row] for row in rows]
    generated = datetime.now(timezone.utc).strftime("%d %b %Y %H:%M UTC")
    if file_format == "csv":
        text = io.StringIO(); writer = csv.writer(text)
        writer.writerow(headers); writer.writerows([[_spreadsheet_cell(value) for value in row] for row in safe_rows])
        payload = io.BytesIO(text.getvalue().encode("utf-8-sig")); payload.seek(0)
        return payload, "text/csv"
    if file_format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook(); sheet = workbook.active; sheet.title = "SAGE Report"
        sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(1,len(headers)))
        sheet.cell(1,1,title).font = Font(size=17,bold=True,color="152536")
        sheet.cell(2,1,subtitle or "Generated " + generated).font = Font(size=11,color="627588")
        sheet.cell(3,1,"Generated " + generated).font = Font(size=10,color="627588")
        header_row = 5
        for index,heading in enumerate(headers,1):
            cell = sheet.cell(header_row,index,heading)
            cell.font = Font(bold=True,color="FFFFFF"); cell.fill = PatternFill("solid",fgColor="203A52")
            cell.alignment = Alignment(vertical="top",wrap_text=True)
        numeric_headers = {"income","expenses","net","amount","value","quantity","total","received","paid","outstanding","requested","approved","actual spent","profit/loss","budget","cash in","cash out","unit value","total value","reorder level","balance","gross pay","deductions","net pay","tax"}
        for r_index,row in enumerate(safe_rows,header_row+1):
            for c_index,value in enumerate(row,1):
                numeric = headers[c_index-1].casefold() in numeric_headers and re.fullmatch(r"[-+]?\d+(?:\.\d+)?",value)
                if numeric and len(value.replace(".","").lstrip("-+0")) <= 15: stored = float(value)
                else: stored = _spreadsheet_cell(value)
                cell = sheet.cell(r_index,c_index,stored); cell.alignment = Alignment(vertical="top",wrap_text=True)
                if numeric and isinstance(stored,float): cell.number_format = '#,##0.00;[Red]-#,##0.00'
                if r_index % 2 == 0: cell.fill = PatternFill("solid",fgColor="F5F8FB")
        sheet.freeze_panes = f"A{header_row+1}"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(max(1,len(headers)))}{max(header_row,header_row+len(safe_rows))}"
        for c_index,heading in enumerate(headers,1):
            lengths = [len(heading)] + [len(row[c_index-1]) for row in safe_rows[:200] if c_index<=len(row)]
            sheet.column_dimensions[get_column_letter(c_index)].width = min(46,max(13,max(lengths,default=13)+2))
        payload = io.BytesIO(); workbook.save(payload); payload.seek(0)
        return payload,"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if file_format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4,landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate,Spacer,LongTable,TableStyle,Paragraph
        font,bold = "Helvetica","Helvetica-Bold"
        regular_path,bold_path = Path("C:/Windows/Fonts/segoeui.ttf"),Path("C:/Windows/Fonts/segoeuib.ttf")
        if regular_path.is_file() and bold_path.is_file():
            if "SageReport" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("SageReport",str(regular_path)))
                pdfmetrics.registerFont(TTFont("SageReportBold",str(bold_path)))
            font,bold = "SageReport","SageReportBold"
        payload = io.BytesIO(); page = landscape(A4) if len(headers)>5 else A4
        doc = SimpleDocTemplate(payload,pagesize=page,rightMargin=25,leftMargin=25,topMargin=28,bottomMargin=35,title=title,author="SAGE")
        title_style = ParagraphStyle("SageTitle",fontName=bold,fontSize=18,leading=24,textColor=colors.HexColor("#172B3D"),spaceAfter=7)
        meta_style = ParagraphStyle("SageMeta",fontName=font,fontSize=9,leading=13,textColor=colors.HexColor("#607487"))
        cell_style = ParagraphStyle("SageCell",fontName=font,fontSize=7.5,leading=10.5,textColor=colors.HexColor("#26394B"),splitLongWords=True)
        header_style = ParagraphStyle("SageHeader",parent=cell_style,fontName=bold,textColor=colors.white)
        def paragraph(value,style=cell_style):
            text = _cell(value).replace("\u2011","-").replace("\u2013","-").replace("\u2014","-")
            return Paragraph(escape(text).replace("\n","<br/>"),style)
        story = [paragraph("SAGE | " + title,title_style),paragraph(subtitle or "Saved business records",meta_style),paragraph("Generated " + generated,meta_style),Spacer(1,17)]
        weights = []
        for heading in headers:
            text = heading.casefold()
            weights.append(3.4 if text in {"before","after","description","detail","latest prices"} else 1.8 if any(word in text for word in ("supplier","requester","department","purchased","approved by","customer")) else 1.4 if "date" in text or "reference" in text else 1.1)
        widths = [doc.width*weight/sum(weights) for weight in weights]
        data = [[paragraph(value,header_style) for value in headers]] + [[paragraph(value) for value in row] for row in safe_rows]
        if not safe_rows: data.append([paragraph("No saved records in this scope.")]+[paragraph("") for _ in headers[1:]])
        table = LongTable(data,colWidths=widths,repeatRows=1,splitInRow=1,hAlign="LEFT")
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#203A52")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F4F7FA")]),("LINEBELOW",(0,0),(-1,0),0.7,colors.HexColor("#203A52")),("GRID",(0,1),(-1,-1),0.3,colors.HexColor("#DDE5EC")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)]))
        story.append(table)
        def footer(canvas,document):
            canvas.saveState(); canvas.setFont(font,8); canvas.setFillColor(colors.HexColor("#728496")); canvas.drawString(25,18,"SAGE - saved business records"); canvas.drawRightString(page[0]-25,18,f"Page {document.page}"); canvas.restoreState()
        doc.build(story,onFirstPage=footer,onLaterPages=footer); payload.seek(0)
        return payload,"application/pdf"
    raise ValueError("Unsupported export format.")
